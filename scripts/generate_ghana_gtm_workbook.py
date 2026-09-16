"""Create the Ghana fabrication-SaaS prospecting workbook.

Run from repo root: python3 scripts/generate_ghana_gtm_workbook.py
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.worksheet.datavalidation import DataValidation


OUTPUT = Path("docs/GHANA-FABRICATION-SaaS-PROSPECT-TRACKER.xlsx")
TODAY = "2026-08-21"

NAVY = "102A43"
BLUE = "007C91"
TEAL = "00A6A6"
PALE = "E8F3F5"
GOLD = "F2B134"
RED = "FDE8E7"
GREEN = "E6F4EA"
WHITE = "FFFFFF"


def add_sheet(wb, name, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 33
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)
    return ws


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "START HERE"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H1")
    ws["A1"] = "VELOXA | Ghana Fabrication SaaS Prospect Tracker"
    ws["A1"].font = Font(bold=True, size=18, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    intro = [
        ("Purpose", "Win the first 5 paid Ghanaian aluminium, glass and façade-fabrication customers before expanding regionally."),
        ("Use this first", "1. Confirm target fit. 2. Contact a named decision maker or the official company channel. 3. Log every touch in OUTREACH LOG. 4. Move the account in PIPELINE. 5. Never send bulk email; personalise from the verified trigger."),
        ("Offer", "A Ghana-native operating system from enquiry and design to GHS quotation, fabrication documents, stock/cutting optimisation, job tracking, QA, dispatch and WhatsApp updates."),
        ("Pilot ask", "A 30-day paid discovery + one live product/workflow pilot. Success condition: one team can create a real quote, factory pack and job workflow without re-keying data."),
        ("Contact standard", "Only public, business-facing details are included. A name is a routing cue, not permission to use a personal email/phone. Verify role and preference before sending."),
        ("Research date", TODAY + ". Recheck website and LinkedIn details immediately before outreach."),
    ]
    for r, (label, value) in enumerate(intro, start=3):
        ws.cell(r, 1, label).font = Font(bold=True, color=NAVY)
        ws.cell(r, 2, value).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 42
    ws["A11"] = "Workbook tabs"
    ws["A11"].font = Font(bold=True, color=WHITE)
    ws["A11"].fill = PatternFill("solid", fgColor=BLUE)
    tabs = [
        ("TARGET ACCOUNTS", "Qualification and buying hypothesis for each company."),
        ("CONTACTS", "Named public contacts and official company channels."),
        ("PIPELINE", "One row per account; your weekly operating view."),
        ("OUTREACH LOG", "Every call, WhatsApp, email, LinkedIn touch and next action."),
        ("EMAIL TEMPLATES", "Personalised email/WhatsApp/call scripts—not mass mail."),
        ("RESEARCH SOURCES", "Public evidence and URLs behind the qualification."),
    ]
    for r, (tab, use) in enumerate(tabs, start=12):
        ws.cell(r, 1, tab).font = Font(bold=True, color=TEAL)
        ws.cell(r, 2, use)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    ws.column_dimensions["A"].width = 23
    ws.column_dimensions["B"].width = 22
    for col in "CDEFGH": ws.column_dimensions[col].width = 16

    targets = [
        ["P1", "Prestodea Ghana Ltd", "Accra + Kumasi + Sunyani", "Architectural & metal fabrication", "Windows/doors; structural & frameless glass; curtain walls; skylights; balustrades; ACP", "10 years; state-of-the-art manufacturing claim; multi-city showrooms; nationwide work", "GHS quote accuracy; multi-branch job visibility; workshop-to-site handoff; material/cutting control", "CEO-led sale; prove one workflow; current system unknown", 91, "Tier 1", "Paid discovery + one product-family pilot", "CEO / Operations", "Samuel Peprah, CEO (publicly listed)", "Official email + phone; LinkedIn/company site", "Website confirms fabrication breadth, multi-city footprint, CEO and company contact."],
        ["P2", "Royal Aluminium Systems Ltd", "Accra", "Aluminium fabricator & installer", "Windows; doors; partitions; false ceilings; curtain walls", "Established 1995; 2,500 sqm workshop; 600 sqm office; West Africa expansion stated", "Factory execution; quote-to-cutting traceability; project stages; management reporting", "Likely mature processes/possible legacy tools; enterprise sales cycle", 90, "Tier 1", "Executive workflow audit + factory-line pilot", "Managing Director / CFO / Operations", "Ernest Taricone, Managing Director; Ajay Kumar Jena, CFO (publicly listed)", "Official site / board page", "Large workshop and regional ambition create clear ROI, but expect longer evaluation."],
        ["P3", "K3 Aluproducts Ltd", "Accra; Nigeria operations", "Glazing & façade contractor", "Windows/doors; curtain wall; ventilated façade; ACP; steel structures", "112 completed projects; Ghana production facility; Nigeria subsidiary; major commercial portfolio", "Complex façade project control; design-to-BOM; procurement; site/install milestones; cross-border visibility", "May have sophisticated internal processes; must not pitch as basic CRM", 89, "Tier 1", "Complex-project control demo using a façade + windows job", "Managing Director / Technical / Operations", "No current named decision maker verified—route through official business channel", "info@k3gh.com; +233 (0) 302-7833-20", "Very strong fit; lead with fewer handoffs and better margin/waste control on complex jobs."],
        ["P4", "E-Facade Ghana Ltd", "Tema / Accra", "Façade and architectural metal contractor", "Curtain walls; ACP; windows/doors; sunshades; architectural metal", "Founded 2013; 11–50 employees (LinkedIn); Ghana and Nigeria projects claimed", "Technical drawings; commercial installed-price estimates; fabrication/install control; QA and traceability", "Need to confirm whether they fabricate in-house and decision process", 86, "Tier 1", "Technical/commercial discovery with a live curtain-wall workflow", "Founder/Director / Technical / Commercial", "Odunayo Odulate and Odulate Philips listed as employees; titles not verified", "info@efacadesghana.com; +233 267 251 830 / +233 507 035 773", "Strong product fit. Treat names as LinkedIn-routing leads only until role is confirmed."],
        ["P5", "Glorious Fabrication Ghana", "Accra", "Glass & aluminium fabrication", "Curtain walls; aluminium windows/doors; stainless balustrades; frameless doors", "Founded 2019; founder-led; residential and commercial focus", "Avoid glass-order errors; instant GHS estimates; hardware/BOM packs; production and installation tracking", "Smaller-company budget sensitivity; sell ROI and a staged subscription", 82, "Tier 1", "Founder pilot focused on frameless/doors plus work-order pack", "Founder & CEO", "Evans Annor Teye, Founder & CEO (publicly listed)", "info@gloriousfabricationgh.com; +233 24 276 3603", "Excellent early-adopter profile: founder directly named and broad enough product mix."],
        ["P6", "Oyiboat Glazing Enterprise", "Accra", "uPVC & aluminium windows/doors fabricator", "uPVC and aluminium doors/windows; staircases; banks, shops, homes, offices, churches, hospitals, estates", "Founded 2015; 11–50 employees (LinkedIn); claims state-of-the-art machinery", "Quote-to-production consistency; custom-measurement capture; capacity/installation scheduling", "Public email/phone not verified in this pass; acquire through company LinkedIn or site first", 78, "Tier 2", "Measurement-to-quote workflow demo", "Owner / Operations / Sales", "Abigail Frempong listed as employee; role not verified", "Company LinkedIn page; address listed on LinkedIn", "Good operational fit; contact discovery required before sequence begins."],
        ["P7", "High-Tech Glazing Ltd", "Greater Accra; serves Kumasi and Takoradi", "uPVC, aluminium & glass fabricator", "Sliding/casement/hinged doors/windows; curtain wall; frameless glass; balustrades", "Established 2016; broad product range; public business phone", "Standardised quoting; hardware and material schedules; multi-region job tracking", "Company size field on LinkedIn looks unreliable; no official email verified", 76, "Tier 2", "WhatsApp-first discovery + standard-product quoting pilot", "Owner / Managing Director / Operations", "Awarkano John Kofi listed as employee; role not verified", "+233 54 615 7895 (public business phone); company LinkedIn", "Useful SaaS-fit profile; qualify budget and project volume on first call."],
    ]
    ws = add_sheet(wb, "TARGET ACCOUNTS", ["ID", "Company", "Geography", "Type", "Publicly evidenced work", "Scale signals", "Likely urgent problem", "Risk / disqualifier", "Fit score /100", "Tier", "Recommended first offer", "Buying roles", "Named route", "Official channel", "Qualification evidence"], targets, [9, 26, 20, 25, 42, 38, 40, 35, 13, 10, 33, 26, 43, 35, 43])
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 9).number_format = "0"
        ws.cell(row, 10).fill = PatternFill("solid", fgColor=GREEN if ws.cell(row, 10).value == "Tier 1" else "FFF4D6")

    contacts = [
        ["C1", "P1", "Prestodea Ghana Ltd", "Samuel Peprah", "CEO / Founder", "Direct decision maker", "prestodeaghana@gmail.com", "+233 244 269 889", "Official company site", "Email first, then call office; ask for a 20-minute operations discovery."],
        ["C2", "P2", "Royal Aluminium Systems Ltd", "Ernest Taricone", "Managing Director", "Direct decision maker", "", "", "Official board page", "Use company contact route and request introduction to MD; personalise around workshop and regional growth."],
        ["C3", "P2", "Royal Aluminium Systems Ltd", "Ajay Kumar Jena", "Chief Financial Officer", "Economic buyer", "", "", "Official board page", "Ask finance/operations question: quote margin, waste, and project cash gates."],
        ["C4", "P3", "K3 Aluproducts Ltd", "", "Managing Director / Technical lead", "Role to identify", "info@k3gh.com", "+233 (0) 302-7833-20", "Official contact page", "Use official channel. Ask who owns fabrication operations and project controls; do not guess a person."],
        ["C5", "P4", "E-Facade Ghana Ltd", "Odunayo Odulate", "Role unverified", "LinkedIn routing lead", "info@efacadesghana.com", "+233 267 251 830", "LinkedIn company page + official contact page", "Contact company email/phone; mention the person only after their role is confirmed."],
        ["C6", "P4", "E-Facade Ghana Ltd", "Odulate Philips", "Role unverified", "LinkedIn routing lead", "info@efacadesghana.com", "+233 507 035 773", "LinkedIn company page + official contact page", "Same approach: verify title and remit before personalising."],
        ["C7", "P5", "Glorious Fabrication Ghana", "Evans Annor Teye", "Founder & CEO", "Direct decision maker", "info@gloriousfabricationgh.com", "+233 24 276 3603", "Official about page", "Founder-to-founder email; lead with preventing glass-order remakes and faster quotes."],
        ["C8", "P6", "Oyiboat Glazing Enterprise", "Abigail Frempong", "Role unverified", "LinkedIn routing lead", "", "", "LinkedIn company page", "Connect through company page; ask who owns estimating/factory operations."],
        ["C9", "P7", "High-Tech Glazing Ltd", "Awarkano John Kofi", "Role unverified", "LinkedIn routing lead", "", "+233 54 615 7895", "LinkedIn company page", "Call public business number; request owner/operations contact. Do not presume person title."],
    ]
    ws = add_sheet(wb, "CONTACTS", ["Contact ID", "Account ID", "Company", "Name", "Title / status", "Contact role", "Public business email", "Public business phone", "Source", "Approach"], contacts, [12, 12, 27, 25, 22, 20, 32, 23, 38, 48])

    pipeline = [[t[0], t[1], t[9], "Not started", "", "", "", "", "", "", "", "", ""] for t in targets]
    ws = add_sheet(wb, "PIPELINE", ["Account ID", "Company", "Tier", "Stage", "Primary owner", "Next action", "Next action date", "Last touch date", "Meeting date", "Pilot hypothesis", "Opportunity value (GHS)", "Probability %", "Weighted value (GHS)"], pipeline, [12, 28, 10, 18, 18, 42, 18, 18, 18, 40, 23, 15, 22])
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 13, f"=IF(OR(K{r}=\"\",L{r}=\"\"),\"\",K{r}*L{r})")
        ws.cell(r, 11).number_format = '"GHS" #,##0'
        ws.cell(r, 12).number_format = "0%"
        ws.cell(r, 13).number_format = '"GHS" #,##0'
    stage_dv = DataValidation(type="list", formula1='"Not started,Researching,Contacted,Replied,Discovery booked,Qualified,Pilot proposed,Pilot active,Commercial negotiation,Won,Nurture,Lost"')
    ws.add_data_validation(stage_dv); stage_dv.add(f"D2:D{ws.max_row}")
    ws.conditional_formatting.add(f"D2:D{ws.max_row}", FormulaRule(formula=['D2="Won"'], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add(f"D2:D{ws.max_row}", FormulaRule(formula=['D2="Lost"'], fill=PatternFill("solid", fgColor=RED)))

    outreach = [["", "", "", "", "", "", "", "", "", "", "", "", ""] for _ in range(25)]
    ws = add_sheet(wb, "OUTREACH LOG", ["Date", "Account ID", "Company", "Contact", "Channel", "Message/template", "Personalisation used", "Outcome", "Objection / insight", "Next action", "Next action date", "Owner", "Consent / preference"], outreach, [15, 12, 26, 24, 15, 24, 37, 25, 35, 35, 18, 18, 22])
    channel_dv = DataValidation(type="list", formula1='"Email,Phone,WhatsApp,LinkedIn,In-person,Referral"')
    ws.add_data_validation(channel_dv); channel_dv.add("E2:E200")

    templates = [
        ["T1", "Email—CEO/founder", "Subject: Could we reduce rework between quote, workshop and installation at [Company]?", "Hi [Name],\n\nI noticed [specific public evidence: e.g., your curtain-wall/window/frameless work]. We are building a Ghana-native operating system for aluminium and glass fabricators: one job record drives the GHS quote, material/hardware schedule, fabrication pack, cutting plan and site progress.\n\nThis is not a generic CRM. We want to understand how [Company] currently prevents price, measurement and fabrication handoff errors on [their product type]. Would you be open to a 20-minute discovery next week? If there is no fit, I will not keep following up.\n\n[Your name]\nVeloxa Technology Limited", "Use one specific project/product claim. Never claim they have a problem you have not verified."],
        ["T2", "Email—operations/technical", "Subject: One drawing → quote → factory pack for [Company]'s [product type] work", "Hi [Name],\n\nWe have built a workflow that turns a design/measurement into a GHS quote, glass/profile/hardware schedule, cutting optimisation and controlled job stages—without retyping the same information across tools.\n\nFor a contractor doing [specific public work], the first question is whether this reduces rework enough to justify a paid pilot. Could I show a 15-minute example based on a [curtain wall / shopfront / window] job?\n\n[Your name]", "For K3/E-Facade lead with complex projects, handoffs and margin control—not basic lead tracking."],
        ["T3", "WhatsApp—first permission ask", "", "Good afternoon [Name]. I am [Name] from Veloxa Technology. We build software for Ghanaian aluminium and glass fabricators—quote, factory pack, cutting/materials and job tracking in one workflow. I saw [specific public evidence]. May I send a one-page overview or book a 15-minute call?", "Use only a public business number and stop if they decline or do not wish to receive messages."],
        ["T4", "Call opener", "", "Hello, my name is [Name] from Veloxa Technology. We work with aluminium and glass fabrication workflows. I am not calling to sell a generic app. Could you point me to the person who owns estimating, factory operations or project delivery? We have a short pilot concept to reduce re-entry between quotation, fabrication documents and site delivery.", "Objective is a warm routing and discovery, not a feature dump."],
        ["T5", "Follow-up after 5 business days", "Subject: Re: [original subject]", "Hi [Name],\n\nClosing the loop on my note. We are looking for one Ghanaian fabrication team to pressure-test a paid pilot around [their product mix]. The pilot starts with a real job and has clear success measures: faster quote turnaround, fewer re-keyed handoffs and a factory-ready pack.\n\nWorth a 20-minute conversation, or should I speak with someone else?\n\n[Your name]", "Send once. Then move to Nurture unless there is engagement."],
    ]
    ws = add_sheet(wb, "EMAIL TEMPLATES", ["ID", "Use", "Subject / opener", "Copy", "Personalisation / guardrail"], templates, [10, 24, 52, 95, 52])

    sources = [
        ["S1", "Prestodea Ghana", "Company scope, CEO, contacts, multi-city presence", "https://www.prestodeagh.com/en/", TODAY],
        ["S2", "Royal Aluminium Systems", "Company scope, history, board, workshop scale", "https://www.royalalu.com/english/index-1.php", TODAY],
        ["S3", "K3 Aluproducts", "Scope, projects, production facility, Nigeria operation", "https://www.k3gh.com/", TODAY],
        ["S4", "K3 Aluproducts", "Official email, phone, Accra office", "https://www.k3gh.com/contact-us/", TODAY],
        ["S5", "E-Facade Ghana", "Scope, services, public contacts", "https://efacadesghana.com/contact/", TODAY],
        ["S6", "E-Facade Ghana LinkedIn", "Employee names, company scale/category", "https://gh.linkedin.com/company/e-facade-ghana-limited", TODAY],
        ["S7", "Glorious Fabrication", "Founder/CEO, services, public contacts", "https://gloriousfabricationgh.com/about/", TODAY],
        ["S8", "Oyiboat Glazing LinkedIn", "Scope, location, scale, employee routing lead", "https://gh.linkedin.com/company/oyiboat-glazing-enterprise", TODAY],
        ["S9", "High-Tech Glazing LinkedIn", "Scope, public business phone, employee routing lead", "https://gh.linkedin.com/company/high-tech-glazing-ltd", TODAY],
        ["S10", "SOFAAMY repo evidence", "Product capabilities and market-fit reference for this playbook", "docs/PROJECT_OUTLINE.md; docs/INDUSTRY-REPORT.md", TODAY],
    ]
    ws = add_sheet(wb, "RESEARCH SOURCES", ["ID", "Company", "What it supports", "URL / repository reference", "Checked"], sources, [10, 28, 45, 75, 15])
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, 4)
        if str(cell.value).startswith("http"):
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"

    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_view.zoomScale = 85
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
